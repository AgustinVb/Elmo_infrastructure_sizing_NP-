"""Arma los datos multianio de DCH (carga on board, caso red) sobre la plantilla
de la mina modelo v4.

Uso (desde la raiz del repo):
    python data/Resultados_finales_tesis/DCH/armar_datos_dch.py

Genera data/Resultados_finales_tesis/DCH/Bateria_482/Costo_fijo/{05C,1C,2C}/P_red
con elmo_data.xlsx, time_series.xlsx y electric_routes_within_time.json (WP2).

Fuentes:
- Plantilla: Mina_modelo_v4/Produccion_base/zona_norte/05C/P_red (caso red:
  Generators con g_max = 0 y Storage con h_max = 0; MarginalCost Profile_fixed,
  0,1 USD/kWh; GenProfiles sin uso porque no hay generacion).
- Por tasa de carga (0,5C / 1C / 2C): hoja chargers y BatteryDegradation de la
  mina modelo v4 zona_norte/{05C,1C,2C}/P_red_gen_bat (cargador de 241 / 482 /
  964 kW a 1.000 USD/kW, gamma 0,0188 / 0,0273 / 0,0451).
- Plan minero DCH 2035-2045 (11 anios, 3 macrobloques): ExtractionGoal,
  NodeAssignment, StationAssignment y Shifts de data/DCH/640kW_2dias/time_series.xlsx
  y FleetByYear de su elmo_data.xlsx. La meta y la asignacion se corrigen para que
  cada LHD activo extraiga solo su calle (ver extraccion_por_calle). FleetByYear lleva
  ademas la columna opera_en_colacion (1 en ANIOS_OPERA_EN_COLACION, hoy ninguno).
- Campania DCH septiembre (rama carga_on_board, Bateria_482/Costo_fijo): hoja
  extraction_nodes (288 puntos con distancias) y los parametros del LH518B
  (peso 54,8 t, velocidad 16,2 km/h).
- Limites fisicos DCH: 3 estaciones, 2 bahias, 1 cargador por bahia,
  p_max_ssee 1.500 kW y p_peak_dist 5.000 kW; costos unitarios de la mina modelo.
- WP2: data/DCH/640kW_2dias/electric_routes_within_time.json (mismo archivo que
  el de la campania de septiembre). Correr con --consumption_model wp2.

Correr siempre con --n_years 11 (ANIOS_HORIZONTE de setup.py es 14 y el plan
trae 11) y --days_per_year 4.
"""
import io
import os
import shutil
import subprocess

import openpyxl
import pandas as pd

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..'))
MM = os.path.join(RAIZ, 'data', 'Resultados_finales_tesis', 'Mina_modelo_v4',
                  'Produccion_base', 'zona_norte')
PLANTILLA = os.path.join(MM, '05C', 'P_red')
DCH_PLAN = os.path.join(RAIZ, 'data', 'DCH', '640kW_2dias')
SEPTIEMBRE = ('carga_on_board:data/Escenarios_DCH_septiembre/Bateria_482/Costo_fijo/'
              'Carga_on_board_fixed_3estaciones_P160kW/elmo_data.xlsx')
SALIDA = os.path.join(RAIZ, 'data', 'Resultados_finales_tesis', 'DCH', 'Bateria_482',
                      'Costo_fijo')

N_ANIOS = 11
TASAS = ['05C', '1C', '2C']
N_ESTACIONES = 3
MAX_BAYS = 2
MAX_CHARGERS_PER_BAY = 1
P_MAX_SSEE = 1500
P_PEAK_DIST = 5000
# Años en que los LHD pueden operar durante la colacion DCH (columna opera_en_colacion de
# FleetByYear). Vacio: se respeta la colacion DCH todos los años. Con 1C y 2C el año 7 es
# factible sin operar en colacion; con 0,5C es infactible igual (falta carga, no tiempo).
ANIOS_OPERA_EN_COLACION = set()


def leer_septiembre():
    crudo = subprocess.run(['git', 'show', SEPTIEMBRE], cwd=RAIZ, check=True,
                           capture_output=True).stdout
    return pd.ExcelFile(io.BytesIO(crudo))


def reescribir(ws, filas, desde):
    """Borra desde la fila `desde` hacia abajo y escribe `filas` ahi."""
    if ws.max_row >= desde:
        ws.delete_rows(desde, ws.max_row - desde + 1)
    for i, fila in enumerate(filas):
        for j, valor in enumerate(fila):
            ws.cell(row=desde + i, column=j + 1, value=valor)


def encabezado(ws):
    return [c.value for c in ws[1]]


def armar_elmo_data(destino, tasa, sep):
    wb = openpyxl.load_workbook(os.path.join(PLANTILLA, 'elmo_data.xlsx'))
    fuente_tasa = openpyxl.load_workbook(os.path.join(MM, tasa, 'P_red_gen_bat', 'elmo_data.xlsx'))

    # LHD: 12 equipos con la fila de la plantilla y los parametros DCH de septiembre.
    ws = wb['LHD']
    cols = encabezado(ws)
    base = [c.value for c in ws[3]]
    lhd_sep = sep.parse('LHD').iloc[1:].iloc[0]
    n_lhd = len(sep.parse('LHD').iloc[1:].dropna(subset=['name']))
    filas = []
    for i in range(1, n_lhd + 1):
        fila = dict(zip(cols, base))
        fila.update(id=i, name=f'LH518B_{i}', model=lhd_sep['model'],
                    weight=float(lhd_sep['weight']), speed=float(lhd_sep['speed']))
        filas.append([fila[c] for c in cols])
    reescribir(ws, filas, 3)

    # stations: costos de la mina modelo, limites fisicos DCH.
    ws = wb['stations']
    cols = encabezado(ws)
    base = dict(zip(cols, [c.value for c in ws[3]]))
    filas = []
    for k in range(1, N_ESTACIONES + 1):
        fila = dict(base)
        fila.update(id=k, station_name=f'station_{k}', max_bays=MAX_BAYS,
                    max_chargers_per_bay=MAX_CHARGERS_PER_BAY, p_max_ssee=P_MAX_SSEE)
        filas.append([fila[c] for c in cols])
    reescribir(ws, filas, 3)

    # chargers: potencia y costo de la tasa, tope de red DCH.
    ws = wb['chargers']
    cols = encabezado(ws)
    ch = dict(zip(encabezado(fuente_tasa['chargers']), [c.value for c in fuente_tasa['chargers'][3]]))
    ch['p_peak_dist'] = P_PEAK_DIST
    reescribir(ws, [[ch[c] for c in cols]], 3)

    # extraction_nodes: los 288 puntos de la campania de septiembre.
    ws = wb['extraction_nodes']
    cols = encabezado(ws)
    nodos = sep.parse('extraction_nodes').iloc[1:]
    reescribir(ws, nodos[cols].values.tolist(), 3)

    # BatteryDegradation de la tasa (mina modelo).
    ws = wb['BatteryDegradation']
    reescribir(ws, [[c.value for c in fuente_tasa['BatteryDegradation'][3]]], 3)

    # FleetByYear del plan DCH, con la columna opera_en_colacion.
    flota = pd.read_excel(os.path.join(DCH_PLAN, 'elmo_data.xlsx'), 'FleetByYear').iloc[1:]
    flota = flota[flota['year'].between(1, N_ANIOS)]
    ws = wb['FleetByYear']
    ws.cell(row=1, column=4, value='opera_en_colacion')
    ws.cell(row=2, column=4, value='bool')
    reescribir(ws, [[i, int(r['year']), int(r['n_active_lhd']), int(int(r['year']) in ANIOS_OPERA_EN_COLACION)]
                    for i, (_, r) in enumerate(flota.iterrows(), 1)], 3)

    wb.save(os.path.join(destino, 'elmo_data.xlsx'))


def extraccion_por_calle(na, eg, sa, sep, anios):
    """Correccion de la usuaria (2026-10-29): en cada año y macrobloque no puede haber mas
    calles con puntos activos que LHD activos. Cada LHD k extrae solo su calle (street_k
    de la campania de septiembre), la produccion del macrobloque se reparte pareja entre las
    calles de sus LHD activos y, dentro de cada calle, pareja entre sus puntos.

    El plan original reparte la meta del macrobloque entre TODOS sus puntos (94/100/94) y,
    con menos LHD que calles, le da dos calles a un LHD (año 2: L3 con 6.782 t/dia). Donde el
    macrobloque tiene sus 4 LHD (años 4-7) no cambia nada. El total por macrobloque y año, y
    los LHD activos de cada año (FleetByYear), son los del plan.

    Devuelve (filas de NodeAssignment, {nodo: {año: meta}})."""
    nodos = sep.parse('extraction_nodes').iloc[1:]
    calle = {r['node_name']: int(str(r['street']).split('_')[1]) for _, r in nodos.iterrows()}
    estacion = dict(zip(sa['elh_name'], sa['station_name']))
    lhd_de_calle = {k: f'LH518B_{k}' for k in set(calle.values())}
    mb_de_nodo = {j: estacion[lhd_de_calle[k]] for j, k in calle.items()}
    puntos_calle = {k: [j for j in eg['name'] if calle[j] == k] for k in lhd_de_calle}

    meta = {j: {y: 0.0 for y in anios} for j in eg['name']}
    activo = {}
    for y in anios:
        asignados = na[na[y] == 1]
        # el plan asigna cada punto a un LHD de su propio macrobloque
        assert all(estacion[l] == mb_de_nodo[j] for l, j in zip(asignados['elhd_name'], asignados['ex_node_name']))
        lhd_activos = set(asignados['elhd_name'])
        activo[y] = lhd_activos
        for m in sorted(set(estacion.values())):
            total = float(eg.loc[eg['name'].map(mb_de_nodo) == m, y].sum())
            lhd_m = sorted(l for l in lhd_activos if estacion[l] == m)
            if total <= 0:
                continue
            assert lhd_m, f'año {y}: {m} tiene meta {total} sin LHD activos'
            for l in lhd_m:
                k = int(l.split('_')[1])
                for j in puntos_calle[k]:
                    meta[j][y] = total / len(lhd_m) / len(puntos_calle[k])

    filas_na = []
    for i, j in enumerate(eg['name'], 1):
        l = lhd_de_calle[calle[j]]
        filas_na.append([i, l, j] + [int(l in activo[y] and meta[j][y] > 0) for y in anios])
    return filas_na, meta


def armar_time_series(destino, sep):
    wb = openpyxl.load_workbook(os.path.join(PLANTILLA, 'time_series.xlsx'))
    ts_dch = os.path.join(DCH_PLAN, 'time_series.xlsx')
    anios = list(range(1, N_ANIOS + 1))

    sh = pd.read_excel(ts_dch, 'Shifts').dropna(subset=['name'])
    cols = encabezado(wb['Shifts'])
    reescribir(wb['Shifts'], [cols] + [[None if pd.isna(r[c]) else r[c] for c in cols]
                                       for _, r in sh.iterrows()], 1)

    na = pd.read_excel(ts_dch, 'NodeAssignment')
    na = na[na['elhd_name'].astype(str).str.match(r'^LH518B_\d+$')]
    sa = pd.read_excel(ts_dch, 'StationAssignment').dropna(subset=['elh_name'])
    eg = pd.read_excel(ts_dch, 'ExtractionGoal').dropna(subset=['name'])
    eg[anios] = eg[anios].fillna(0)
    filas_na, meta = extraccion_por_calle(na, eg, sa, sep, anios)

    reescribir(wb['NodeAssignment'], [['id', 'elhd_name', 'ex_node_name'] + anios] + filas_na, 1)
    reescribir(wb['StationAssignment'],
               [['id', 'elh_name', 'station_name']]
               + [[i, r['elh_name'], r['station_name']] for i, (_, r) in enumerate(sa.iterrows(), 1)], 1)
    reescribir(wb['ExtractionGoal'],
               [['id', 'name'] + anios]
               + [[i, j] + [meta[j][y] for y in anios] for i, j in enumerate(eg['name'], 1)], 1)

    wb.save(os.path.join(destino, 'time_series.xlsx'))


def main():
    sep = leer_septiembre()
    for tasa in TASAS:
        destino = os.path.join(SALIDA, tasa, 'P_red')
        os.makedirs(destino, exist_ok=True)
        armar_elmo_data(destino, tasa, sep)
        armar_time_series(destino, sep)
        shutil.copy(os.path.join(DCH_PLAN, 'electric_routes_within_time.json'), destino)
        print('listo', os.path.relpath(destino, RAIZ))


if __name__ == '__main__':
    main()
